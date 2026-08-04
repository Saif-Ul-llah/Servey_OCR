"""Confidence and provenance data, kept out of the deliverable.

The reference workbook has three columns and one sheet, and the existing
workflow reads it directly, so nothing extra may be added to it. Everything the
operator and the debugging process need lives here instead, in a file named
after the deliverable with a ``_review`` suffix.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from local_ocr.assembly.audit import describe
from local_ocr.models import BatchAudit, OutputRow, PageResult, RowRecord, Status

REVIEW_HEADERS = (
    "Page",
    "Row",
    "Column",
    "Value",
    "RawOCR",
    "Confidence",
    "Status",
    "AppliedRules",
    "Alternatives",
    "EngineOutputs",
)

SUMMARY_HEADERS = ("Metric", "Value")

#: Red for "must be looked at", amber for "a rule changed this".
FILL_REVIEW = PatternFill("solid", fgColor="FFC7CE")
FILL_CORRECTED = PatternFill("solid", fgColor="FFEB9C")


def sidecar_path(deliverable: Path | str, suffix: str = "_review") -> Path:
    deliverable = Path(deliverable)
    return deliverable.with_name(f"{deliverable.stem}{suffix}{deliverable.suffix}")


def _cell_rows(records: list[RowRecord]):
    for record in records:
        for name, cell in (
            ("survey", record.survey),
            ("meter", record.meter),
            ("remarks", record.remarks),
        ):
            if cell.is_empty and cell.status in (Status.EMPTY, Status.OK):
                continue
            yield (
                record.page,
                record.index,
                name,
                cell.text,
                cell.raw_text,
                round(cell.confidence, 3),
                cell.status.value,
                ", ".join(cell.applied_rules),
                " | ".join(cell.candidates),
                " | ".join(f"{k}={v}" for k, v in cell.engine_outputs.items()),
            )


def write_sidecar(
    path: Path | str,
    records: list[RowRecord],
    rows: list[OutputRow],
    audit: BatchAudit,
    pages: list[PageResult] | None = None,
    skipped: list[tuple[RowRecord, str]] | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    review = wb.active
    review.title = "Review"
    review.append(list(REVIEW_HEADERS))
    for values in _cell_rows(records):
        review.append(list(values))
        status = values[6]
        if Status(status).needs_review:
            fill = FILL_REVIEW
        elif status == Status.CORRECTED.value:
            fill = FILL_CORRECTED
        else:
            continue
        for column in range(1, len(REVIEW_HEADERS) + 1):
            review.cell(row=review.max_row, column=column).fill = fill
    review.freeze_panes = "A2"
    _autosize(review)

    summary = wb.create_sheet("Summary")
    summary.append(list(SUMMARY_HEADERS))
    flagged = sum(1 for r in rows if r.status.needs_review)
    corrected = sum(1 for r in rows if r.status is Status.CORRECTED)
    for metric, value in (
        ("Pages processed", len(pages or [])),
        ("Rows exported", len(rows)),
        ("Rows flagged for review", flagged),
        ("Rows auto-corrected", corrected),
        ("Rows skipped (markers / non-Latin)", len(skipped or [])),
        ("Survey gaps", len(audit.survey_gaps)),
        ("Survey repeats", len(audit.survey_repeats)),
        ("Duplicate meter codes", len(audit.duplicate_meters)),
        ("Row-count mismatches", len(audit.row_count_mismatches)),
        ("Structural audit clean", "yes" if audit.clean else "no"),
    ):
        summary.append([metric, value])
    summary.append([])
    summary.append(["Audit findings", ""])
    for line in describe(audit):
        summary.append(["", line])
    if skipped:
        summary.append([])
        summary.append(["Skipped lines", ""])
        for record, reason in skipped:
            summary.append([f"{record.page} row {record.index}", reason])
    _autosize(summary)

    wb.save(path)
    return path


def _autosize(worksheet, limit: int = 60) -> None:
    for column in range(1, worksheet.max_column + 1):
        widest = max(
            (len(str(worksheet.cell(row=r, column=column).value or "")) for r in range(1, worksheet.max_row + 1)),
            default=0,
        )
        worksheet.column_dimensions[get_column_letter(column)].width = min(max(widest + 2, 10), limit)
