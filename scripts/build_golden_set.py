"""Build samples/ground_truth.csv from the reference workbook + page manifest.

The 12 sample photos cover a contiguous run of the reference workbook
(KE-570..KE-683 == rows 596..882), so the ground truth already exists and only
needs slicing per page. samples/page_manifest.yaml holds the mapping.

The emitted CSV is the contract for both scripts/eval_engines.py and the
integration test, and records the *cell type* of each remark, because the
reference workbook stores a lone side code as a number (-37) and a combined one
as text ('-42,-128'). Reproducing that distinction is part of the deliverable.

Usage:
    python scripts/build_golden_set.py
    python scripts/build_golden_set.py --copy-images   # also stage samples/images_lowres/
"""

from __future__ import annotations

import argparse
import csv
import shutil
import sys
from pathlib import Path

import openpyxl
import yaml

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MANIFEST = REPO_ROOT / "samples" / "page_manifest.yaml"
DEFAULT_OUT = REPO_ROOT / "samples" / "ground_truth.csv"

FIELDNAMES = [
    "page_id",
    "row_index",       # 0-based order of the line on that page
    "workbook_row",    # 1-based worksheet row, for traceability
    "survey",          # "KE-646" on a group-first row, "" otherwise
    "meter",
    "remarks",
    "remarks_type",    # "text" | "number" | ""
    "is_group_first",
]


def load_manifest(path: Path) -> dict:
    with path.open(encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def check_tiling(pages: list[dict]) -> None:
    """The page spans must tile the covered range exactly -- no gaps, no overlaps.

    A gap means a photo is missing; an overlap means two photos claim the same
    notebook page. Either one silently corrupts the golden set, so it is fatal.
    """
    ordered = sorted(pages, key=lambda p: p["first_row"])
    for previous, current in zip(ordered, ordered[1:]):
        expected = previous["last_row"] + 1
        if current["first_row"] != expected:
            kind = "gap" if current["first_row"] > expected else "overlap"
            raise SystemExit(
                f"manifest {kind}: {previous['id']} ends at row {previous['last_row']}, "
                f"{current['id']} starts at row {current['first_row']} (expected {expected})"
            )


def remark_type(value: object) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return ""
    return "number" if isinstance(value, (int, float)) else "text"


def build_rows(workbook: Path, sheet: str, pages: list[dict]) -> list[dict]:
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.active

    out: list[dict] = []
    for page in sorted(pages, key=lambda p: p["first_row"]):
        for offset, wb_row in enumerate(range(page["first_row"], page["last_row"] + 1)):
            survey, meter, remarks = (ws.cell(row=wb_row, column=c).value for c in (1, 2, 3))
            out.append(
                {
                    "page_id": page["id"],
                    "row_index": offset,
                    "workbook_row": wb_row,
                    "survey": "" if survey is None else str(survey).strip(),
                    "meter": "" if meter is None else str(meter).strip(),
                    "remarks": "" if remarks is None else str(remarks).strip(),
                    "remarks_type": remark_type(remarks),
                    "is_group_first": "1" if survey is not None else "0",
                }
            )
    wb.close()
    return out


def copy_images(manifest: dict, dest: Path) -> int:
    src_folder = REPO_ROOT / manifest["source_folder"]
    dest.mkdir(parents=True, exist_ok=True)
    copied = 0
    for page in sorted(manifest["pages"], key=lambda p: p["first_row"]):
        src = src_folder / page["source"]
        if not src.exists():
            print(f"  warn: missing source image for {page['id']}: {src.name}", file=sys.stderr)
            continue
        shutil.copy2(src, dest / f"{page['id']}{src.suffix.lower()}")
        copied += 1
    return copied


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument(
        "--copy-images",
        action="store_true",
        help="stage the sample photos into samples/images_lowres/ as page_NN.jpeg",
    )
    args = parser.parse_args(argv)

    manifest = load_manifest(args.manifest)
    pages = manifest["pages"]
    check_tiling(pages)

    workbook = REPO_ROOT / manifest["workbook"]
    if not workbook.exists():
        print(f"error: workbook not found: {workbook}", file=sys.stderr)
        return 1

    rows = build_rows(workbook, manifest["sheet"], pages)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)

    first = min(p["first_row"] for p in pages)
    last = max(p["last_row"] for p in pages)
    print(f"{len(pages)} pages tile workbook rows {first}-{last} with no gaps or overlaps")
    print(f"wrote {len(rows)} ground-truth rows to {args.out.relative_to(REPO_ROOT)}")

    groups = sum(1 for r in rows if r["is_group_first"] == "1")
    remarked = sum(1 for r in rows if r["remarks"])
    numeric = sum(1 for r in rows if r["remarks_type"] == "number")
    print(f"  {groups} survey groups, {remarked} rows with remarks ({numeric} numeric side codes)")

    if args.copy_images:
        dest = REPO_ROOT / "samples" / "images_lowres"
        copied = copy_images(manifest, dest)
        print(f"staged {copied} images into {dest.relative_to(REPO_ROOT)}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
