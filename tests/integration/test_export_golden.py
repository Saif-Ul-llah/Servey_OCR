"""End-to-end check of correction + assembly + export against the real workbook.

The recognizer is simulated as perfect: meter codes are fed in the way they are
written on the page (``SFT 77430``), placeholders as their handwritten wording,
side codes as the marks the surveyor makes. Everything downstream of the OCR
engine then has to reproduce workbook rows 596-882 exactly.

The two page-straddling groups are simulated faithfully: on page_05 and page_11
the surveyor rewrites the survey number at the top of the continuing page, so
the test feeds that repeat in and asserts the assembly stage demotes it instead
of emitting a second ``KE-619`` / ``KE-673``.
"""

from __future__ import annotations

import csv
import re

import openpyxl
import pytest
import yaml

from local_ocr.assembly import assemble
from local_ocr.correction import Corrector
from local_ocr.export import HEADERS, SHEET_NAME, read_workbook, write_workbook
from local_ocr.models import Cell, Column, PageResult, RowRecord
from local_ocr.paths import GROUND_TRUTH, PAGE_MANIFEST, REPO_ROOT

AS_WRITTEN = {
    "NoMeter": "No meter",
    "EmptyPlot": "Empty plot",
    "Empty": "Empty",
    "Plot": "Plot",
    "MeterMissing": "Meter missing",
}

METER_RE = re.compile(r"^([A-Z]{2,3})(\d+)$")
ANNOTATION_RE = re.compile(r"^-\d{1,4}$")

pytestmark = pytest.mark.skipif(
    not GROUND_TRUTH.exists(),
    reason="run scripts/build_golden_set.py first",
)


def as_written(meter: str) -> str:
    if meter in AS_WRITTEN:
        return AS_WRITTEN[meter]
    match = METER_RE.match(meter)
    return f"{match.group(1)} {match.group(2)}" if match else meter


def split_remark(value: str, kind: str) -> tuple[list[str], str]:
    """Recover (side codes, free text) from a workbook REMARKS value."""
    if kind == "number":
        return [str(int(float(value)))], ""
    if not value:
        return [], ""
    annotations: list[str] = []
    rest: list[str] = []
    for part in value.split(","):
        (annotations if ANNOTATION_RE.fullmatch(part.strip()) else rest).append(part.strip() if ANNOTATION_RE.fullmatch(part.strip()) else part)
    return annotations, ",".join(rest).strip()


def as_marked(annotations: list[str]) -> str:
    """Render side codes the way the surveyor writes them: -37 was written +37."""
    return " ".join(f"+{code.lstrip('-')}" for code in annotations)


@pytest.fixture(scope="module")
def manifest() -> dict:
    return yaml.safe_load(PAGE_MANIFEST.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def golden() -> dict[str, list[dict[str, str]]]:
    pages: dict[str, list[dict[str, str]]] = {}
    with GROUND_TRUTH.open(encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            pages.setdefault(row["page_id"], []).append(row)
    return pages


@pytest.fixture(scope="module")
def simulated_pages(manifest, golden) -> list[PageResult]:
    """Build PageResults as a perfect recognizer would have produced them."""
    corrector = Corrector()
    entries = sorted(manifest["pages"], key=lambda p: p["first_row"])
    results: list[PageResult] = []
    last_survey_digits: str | None = None

    for entry in entries:
        rows = golden[entry["id"]]
        records: list[RowRecord] = []

        for position, row in enumerate(rows):
            survey_written = row["survey"].replace("KE-", "")
            if position == 0 and entry["continues_previous_group"]:
                # The surveyor rewrites the number at the top of a continuing
                # page; the workbook does not repeat it.
                assert not survey_written, "a continuing page must not open a new group"
                survey_written = last_survey_digits or ""

            annotations, text = split_remark(row["remarks"], row["remarks_type"])
            remarks_written = " ".join(part for part in (as_marked(annotations), text) if part)

            records.append(
                RowRecord(
                    page=entry["id"],
                    index=position,
                    survey=Cell(column=Column.SURVEY, raw_text=survey_written, confidence=1.0),
                    meter=Cell(column=Column.METER, raw_text=as_written(row["meter"]), confidence=1.0),
                    remarks=Cell(column=Column.REMARKS, raw_text=remarks_written, confidence=1.0),
                )
            )

        corrector.correct_page(records)
        results.append(PageResult(path=entry["id"], rows=records, detected_rules=len(records)))
        last_survey_digits = next(
            (r["survey"].replace("KE-", "") for r in reversed(rows) if r["survey"]),
            last_survey_digits,
        )

    return results


@pytest.fixture(scope="module")
def reference_rows() -> list[tuple]:
    """Rows 596-882 of the reference workbook, values and types as stored."""
    workbook = REPO_ROOT / "Samples or examle data" / "SACHAL SURVEY - IMAGES DATA.xlsx"
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb.active
    rows = [
        tuple(ws.cell(row=r, column=c).value for c in (1, 2, 3))
        for r in range(596, 883)
    ]
    wb.close()
    return rows


def test_row_count_matches(simulated_pages, reference_rows):
    batch = assemble(simulated_pages)
    assert len(batch.rows) == len(reference_rows) == 287


def test_survey_column_matches_exactly(simulated_pages, reference_rows):
    batch = assemble(simulated_pages)
    produced = [row.survey for row in batch.rows]
    expected = [row[0] for row in reference_rows]
    mismatches = [
        (i + 596, e, p) for i, (e, p) in enumerate(zip(expected, produced)) if e != p
    ]
    assert not mismatches, f"survey mismatches at workbook rows: {mismatches[:10]}"


def test_meter_column_matches_exactly(simulated_pages, reference_rows):
    batch = assemble(simulated_pages)
    produced = [row.meter for row in batch.rows]
    expected = [row[1] for row in reference_rows]
    mismatches = [
        (i + 596, e, p) for i, (e, p) in enumerate(zip(expected, produced)) if e != p
    ]
    assert not mismatches, f"meter mismatches at workbook rows: {mismatches[:10]}"


def test_side_codes_keep_their_numeric_cell_type(simulated_pages, reference_rows):
    batch = assemble(simulated_pages)
    for index, (produced, expected) in enumerate(zip(batch.rows, reference_rows)):
        if isinstance(expected[2], (int, float)):
            assert isinstance(produced.remarks, int), (
                f"workbook row {index + 596} stores {expected[2]!r} as a number; "
                f"produced {produced.remarks!r}"
            )
            assert produced.remarks == expected[2]


def test_remarks_match_after_normalisation(simulated_pages, reference_rows):
    """Free text is compared loosely.

    The human transcription normalises as it goes: ``' 4 Shops'`` carries a
    stray leading space at workbook row 817, ``'-37,SHOP'`` is uppercased at row
    811 where the page reads "Shop", and ``'Hair Cutting Saloon'`` is title-cased
    at row 610. REMARKS is therefore held to a case- and whitespace-insensitive
    comparison. Meter and survey are not -- those must match exactly.
    """
    batch = assemble(simulated_pages)
    mismatches = []
    for index, (produced, expected) in enumerate(zip(batch.rows, reference_rows)):
        want = "" if expected[2] is None else str(expected[2]).strip().casefold()
        got = "" if produced.remarks is None else str(produced.remarks).strip().casefold()
        if want != got:
            mismatches.append((index + 596, want, got))
    assert not mismatches, f"remark mismatches: {mismatches[:10]}"


def test_continuing_group_is_not_relabelled(simulated_pages):
    """KE-619 and KE-673 each span two photographs and must appear once."""
    batch = assemble(simulated_pages)
    surveys = [row.survey for row in batch.rows if row.survey]
    assert surveys.count("KE-619") == 1
    assert surveys.count("KE-673") == 1
    assert len(surveys) == len(set(surveys)) == 114


def test_structural_audit_is_clean(simulated_pages):
    batch = assemble(simulated_pages)
    assert batch.audit.survey_gaps == []
    assert batch.audit.survey_repeats == []
    assert batch.audit.duplicate_meters == []
    assert batch.audit.row_count_mismatches == []
    assert batch.audit.clean


def test_written_workbook_reloads_identically(simulated_pages, reference_rows, tmp_path):
    batch = assemble(simulated_pages)
    path = write_workbook(batch.rows, tmp_path / "out.xlsx", backup_previous=False)

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.title == SHEET_NAME
    assert tuple(ws.cell(row=1, column=c).value for c in (1, 2, 3)) == HEADERS
    assert ws.max_row == 288  # header + 287 rows

    for index, expected in enumerate(reference_rows, start=2):
        assert ws.cell(row=index, column=1).value == expected[0]
        assert ws.cell(row=index, column=2).value == expected[1]

    reloaded = read_workbook(path)
    assert len(reloaded) == 287
    wb.close()
